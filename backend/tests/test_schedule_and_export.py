"""Follow-up schedule overrides (per company / per stage) and the Excel export."""
import io
from datetime import datetime, timedelta

import pandas as pd
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.models.lead import AIResponseClassification, EmailStatus, Lead
from app.pipeline import outreach_graph
from app.pipeline.export_leads import COLUMNS, leads_to_xlsx
from app.services import schedule_service
from app.services.schedule_service import Schedule
from app.services.storage_service import get_lead_store


def make_lead(**overrides) -> Lead:
    defaults = dict(created_time=datetime.utcnow(), full_name="Test Lead", email="t@example.com")
    defaults.update(overrides)
    return Lead(**defaults)


def test_default_wait_is_three_days():
    assert Schedule().wait_days_for(make_lead()) == 3


def test_stage_override_applies_to_that_follow_up_only():
    schedule = Schedule(by_stage={"2": 7})
    assert schedule.wait_days_for(make_lead(follow_up_count=0)) == 3  # wait before follow-up 1
    assert schedule.wait_days_for(make_lead(follow_up_count=1)) == 7  # wait before follow-up 2
    assert schedule.wait_days_for(make_lead(follow_up_count=2)) == 3


def test_company_override_beats_stage_override_and_ignores_case():
    schedule = Schedule(by_company={"Acme Industrial": 10}, by_stage={"1": 5})
    acme = make_lead(company="  acme industrial ")
    other = make_lead(company="Reilly Freight")
    assert schedule.wait_days_for(acme) == 10
    assert schedule.wait_days_for(other) == 5


def test_out_of_range_wait_is_rejected():
    with pytest.raises(ValueError):
        Schedule(by_company={"Acme": 0})
    with pytest.raises(ValueError):
        Schedule(default_days=500)


def test_next_follow_up_due_uses_the_saved_schedule():
    schedule_service.set_schedule(Schedule(by_company={"Acme Industrial": 6}))
    sent_at = datetime.utcnow()
    lead = make_lead(company="Acme Industrial", email_status=EmailStatus.sent, last_sent_at=sent_at)
    assert outreach_graph.next_follow_up_due_at(lead) == sent_at + timedelta(days=6)


def test_schedule_api_round_trips():
    client = TestClient(app)
    assert client.get("/pipeline/schedule").json()["default_days"] == 3
    body = {"default_days": 4, "by_company": {"Acme Industrial": 9}, "by_stage": {"1": 2}}
    assert client.put("/pipeline/schedule", json=body).status_code == 200
    assert client.get("/pipeline/schedule").json() == body
    assert client.put("/pipeline/schedule", json={"default_days": 0}).status_code == 422


def test_export_contains_only_the_requested_leads_with_reply():
    store = get_lead_store()
    keep = make_lead(email="keep@example.com", first_name="Keep", last_name="Me", company="Acme Industrial",
                     position="CEO", employee_count="120", location="Stamford, CT",
                     email_status=EmailStatus.sent, ai_response_classification=AIResponseClassification.interested,
                     last_reply_text="Yes, let's talk Thursday", last_reply_at=datetime(2026, 10, 7, 12, 0))
    skip = make_lead(email="skip@example.com", first_name="Skip", last_name="Me")
    store.save(keep)
    store.save(skip)

    resp = TestClient(app).post("/leads/export", json={"lead_ids": [keep.id], "label": "Acme / interested"})
    assert resp.status_code == 200
    assert resp.headers["content-disposition"] == 'attachment; filename="Acme___interested.xlsx"'

    df = pd.read_excel(io.BytesIO(resp.content))
    assert list(df.columns) == COLUMNS
    assert len(df) == 1
    row = df.iloc[0]
    assert row["Name"] == "Test Lead" and row["Email"] == "keep@example.com"
    assert row["Company"] == "Acme Industrial"
    assert row["Reply"] == "Yes, let's talk Thursday"
    assert row["Reply type"] == "Interested"
    assert row["Stage"] == "Interested — awaiting review"
    assert row["Qualification"] == "Not yet reviewed"  # plain words, not the stored value "unset"


def test_export_of_nothing_is_just_the_header_row():
    df = pd.read_excel(io.BytesIO(leads_to_xlsx([])))
    assert list(df.columns) == COLUMNS and df.empty


def _sheet(leads, label=""):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(leads_to_xlsx(leads, label)))


def test_the_sheet_is_laid_out_for_reading():
    wb = _sheet([make_lead(email="a@x.com", full_name="A One", company="Acme"), make_lead(email="b@x.com", full_name="B Two")], "All leads - Acme")
    ws = wb["Leads"]
    assert wb.sheetnames == ["Leads", "About this export"]
    assert ws.freeze_panes == "C2"                                   # header + name + company stay in view
    assert ws.auto_filter.ref == f"A1:{chr(64 + len(COLUMNS))}3"     # filter buttons over header and both rows
    assert ws["A1"].font.b and ws["A1"].fill.fgColor.rgb.endswith("CC1B1B")
    widths = {ws.cell(row=1, column=i + 1).value: ws.column_dimensions[chr(65 + i)].width for i in range(len(COLUMNS))}
    assert widths["Reply"] >= 50 and widths["Notes"] >= 40 and widths["Email"] >= 30   # nothing cramped
    assert ws.cell(row=2, column=COLUMNS.index("Reply") + 1).alignment.wrap_text is True


def test_times_are_shown_in_eastern_not_utc_and_formatted_as_dates():
    # 2026-10-07 16:00 UTC is 12:00 noon Eastern (EDT)
    lead = make_lead(email="t@x.com", last_sent_at=datetime(2026, 10, 7, 16, 0), last_reply_at=datetime(2026, 1, 15, 17, 30))
    ws = _sheet([lead])["Leads"]
    sent = ws.cell(row=2, column=COLUMNS.index("Last email sent (ET)") + 1)
    reply = ws.cell(row=2, column=COLUMNS.index("Reply received (ET)") + 1)
    assert sent.value == datetime(2026, 10, 7, 12, 0) and sent.number_format == "mmm d, yyyy h:mm AM/PM"
    assert reply.value == datetime(2026, 1, 15, 12, 30)  # winter: EST, UTC-5


def test_status_cells_are_coloured_and_the_profile_url_is_a_link():
    lead = make_lead(email="c@x.com", ai_response_classification=AIResponseClassification.interested,
                     qualified_status="qualified", linkedin_url="https://www.linkedin.com/in/x")
    ws = _sheet([lead])["Leads"]
    assert ws.cell(row=2, column=COLUMNS.index("Stage") + 1).fill.fgColor.rgb.endswith("E6F1E1")          # positive
    assert ws.cell(row=2, column=COLUMNS.index("Qualification") + 1).fill.fgColor.rgb.endswith("E6F1E1")
    assert ws.cell(row=2, column=COLUMNS.index("Profile URL") + 1).hyperlink.target == "https://www.linkedin.com/in/x"


def test_the_about_sheet_says_what_was_exported():
    ws = _sheet([make_lead(email="d@x.com")], "Needs review - Acme")["About this export"]
    values = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, 6)}
    assert values["View"] == "Needs review - Acme" and values["Leads included"] == 1
    assert "Eastern" in values["Times"]


def test_an_empty_export_still_has_the_header_and_filters():
    ws = _sheet([])["Leads"]
    assert [c.value for c in ws[1]] == COLUMNS and ws.max_row == 1
