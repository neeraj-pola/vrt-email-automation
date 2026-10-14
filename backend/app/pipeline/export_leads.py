"""
Excel export of a set of leads, laid out to be read by staff: the columns people look at first on the
left (name, company, stage), the long ones (reply, notes) wrapped and wide, headers in the brand colour,
the header row and the first two columns frozen, filter buttons on every column, times in US Eastern
(the company's timezone, not UTC), and plain-language values ("Not yet reviewed", not "unset").
"""
from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Optional
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models.lead import Lead

_ET = ZoneInfo("America/New_York")
_DATE_FORMAT = "mmm d, yyyy h:mm AM/PM"

BRAND_RED = "CC1B1B"
_HEADER_FILL = PatternFill("solid", fgColor=BRAND_RED)
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_THIN = Side(style="thin", color="E3D6CD")
_BORDER = Border(bottom=_THIN)
_ZEBRA = PatternFill("solid", fgColor="FAF6F3")
_TONES = {  # soft fills for the status columns, same meaning as the dashboard's colours
    "positive": PatternFill("solid", fgColor="E6F1E1"),
    "negative": PatternFill("solid", fgColor="FBE3E3"),
    "warning": PatternFill("solid", fgColor="E7EBF6"),
}

# (header, width, wrap text)
_COLUMNS: list[tuple[str, int, bool]] = [
    ("Name", 22, False),
    ("Company", 26, False),
    ("Position", 26, False),
    ("Employees", 11, False),
    ("Location", 22, False),
    ("Email", 32, False),
    ("Stage", 28, False),
    ("Qualification", 17, False),
    ("Reply type", 15, False),
    ("Reply", 52, True),
    ("Reply received (ET)", 21, False),
    ("Requested availability", 24, True),
    ("Scheduled call (ET)", 21, False),
    ("Last email sent (ET)", 21, False),
    ("Follow-ups sent", 10, False),
    ("Email status", 15, False),
    ("Notes", 44, True),
    ("CRM status", 16, False),
    ("Sale status", 16, False),
    ("Profile URL", 34, False),
    ("Connected on (ET)", 21, False),
]
COLUMNS = [c[0] for c in _COLUMNS]

_QUALIFICATION = {"unset": "Not yet reviewed", "qualified": "Qualified", "not_qualified": "Not qualified"}
_EMAIL_STATUS = {"not_sent": "Not sent", "sent": "Sent", "bounced": "Invalid email",
                 "unsubscribed": "Unsubscribed", "exhausted": "No response"}
_REPLY_TYPE = {"interested": "Interested", "not_interested": "Not interested", "unclear": "Unclear"}


def _enum(value) -> Optional[str]:
    return getattr(value, "value", value)


def _et(value: Optional[datetime]) -> Optional[datetime]:
    """Stored times are naive UTC; show US Eastern, as a naive value Excel can format."""
    if value is None:
        return None
    aware = value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value
    return aware.astimezone(_ET).replace(tzinfo=None)


def _stage_label(lead: Lead) -> tuple[str, str]:
    """(label, tone) -- same wording and colour meaning as the dashboard (frontend/src/lib/stage.ts)."""
    status, ai, qual = _enum(lead.email_status), _enum(lead.ai_response_classification), _enum(lead.qualified_status)
    if status == "bounced":
        return "Invalid email", "negative"
    if status == "unsubscribed":
        return "Unsubscribed", "negative"
    if status == "exhausted":
        return "No response (exhausted)", "warning"
    if ai == "interested":
        if lead.scheduled_call_time:
            return "Call scheduled", "positive"
        if qual == "qualified":
            return "Qualified — awaiting call", "positive"
        if qual == "not_qualified":
            return "Interested — not qualified", "positive"
        return "Interested — awaiting review", "positive"
    if ai == "not_interested":
        return "Not interested", "negative"
    if ai == "unclear":
        return "Unclear reply", "warning"
    if status == "not_sent":
        return ("Draft ready" if lead.current_draft_subject else "Not contacted"), ""
    if status == "sent":
        return (f"Follow-up {lead.follow_up_count} sent" if lead.follow_up_count > 0 else "Initial email sent"), ""
    return "Unknown", ""


def _row(lead: Lead) -> tuple[list, str, str]:
    stage, tone = _stage_label(lead)
    qual = _enum(lead.qualified_status)
    ai = _enum(lead.ai_response_classification)
    values = [
        lead.full_name, lead.company, lead.position, lead.employee_count, lead.location, lead.email,
        stage, _QUALIFICATION.get(qual, qual), _REPLY_TYPE.get(ai, "" if ai in (None, "pending") else ai),
        lead.last_reply_text, _et(lead.last_reply_at), lead.requested_availability,
        _et(lead.scheduled_call_time), _et(lead.last_sent_at), lead.follow_up_count,
        _EMAIL_STATUS.get(_enum(lead.email_status), _enum(lead.email_status)),
        lead.notes, lead.crm_status, lead.sale_status, lead.linkedin_url, _et(lead.created_time),
    ]
    qual_tone = {"qualified": "positive", "not_qualified": "negative"}.get(qual, "")
    return values, tone, qual_tone


def leads_to_xlsx(leads: list[Lead], label: str = "") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    for col, (header, width, _wrap) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 32

    stage_col, qual_col = COLUMNS.index("Stage") + 1, COLUMNS.index("Qualification") + 1
    url_col = COLUMNS.index("Profile URL") + 1
    for r, lead in enumerate(leads, start=2):
        values, stage_tone, qual_tone = _row(lead)
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=_COLUMNS[c - 1][2])
            cell.border = _BORDER
            if isinstance(value, datetime):
                cell.number_format = _DATE_FORMAT
            if r % 2 == 1:
                cell.fill = _ZEBRA
        if stage_tone:
            ws.cell(row=r, column=stage_col).fill = _TONES[stage_tone]
        if qual_tone:
            ws.cell(row=r, column=qual_col).fill = _TONES[qual_tone]
        if values[url_col - 1]:
            link = ws.cell(row=r, column=url_col)
            link.hyperlink = values[url_col - 1]
            link.font = Font(color="1F5FBF", underline="single")

    ws.freeze_panes = "C2"  # header row, plus name and company stay in view while scrolling
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_COLUMNS))}{max(1, len(leads) + 1)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    about = wb.create_sheet("About this export")
    about.column_dimensions["A"].width = 18
    about.column_dimensions["B"].width = 70
    rows = [
        ("View", label or "All leads"),
        ("Leads included", len(leads)),
        ("Exported", _et(datetime.now(timezone.utc))),
        ("Times", "US Eastern (the company's timezone)"),
        ("Source", "VRT Management Group lead outreach dashboard"),
    ]
    for i, (key, value) in enumerate(rows, start=1):
        about.cell(row=i, column=1, value=key).font = Font(bold=True)
        cell = about.cell(row=i, column=2, value=value)
        cell.alignment = Alignment(horizontal="left", wrap_text=True)
        if isinstance(value, datetime):
            cell.number_format = _DATE_FORMAT

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
